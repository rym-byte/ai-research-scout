"""AI Research Scout - Web 应用"""
import os
import asyncio
import secrets
import io
import base64
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file, make_response, redirect, url_for
from markupsafe import escape

from src.collectors import HNCollector, WebCollector
from src.analyzers import ResearchSynthesizer
from src.analyzers.groq_analyzer import analyze_with_groq
from src.outputs import MarkdownOutput

# 尝试导入 Word/PDF 导出库
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# 存储最新报告（用于 Web UI 展示）
latest_report = {
    'topic': '',
    'report': None,
    'groq_analysis': None,
    'items': [],
    'created_at': None
}

app = Flask(__name__)
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', 'ai-research-scout-2026')

# API Token 用于定时任务调用
API_TOKEN = os.environ.get('API_TOKEN', 'sk_live_' + secrets.token_hex(16))

# 存储历史记录
history = []


async def run_research(topic: str, sources: list, limit: int):
    """执行研究"""
    all_items = []
    collectors = []

    if 'hn' in sources:
        collectors.append(('Hacker News', HNCollector()))
    if 'web' in sources:
        collectors.append(('Web', WebCollector()))

    results = {'sources': {}, 'items': [], 'report': None}

    for name, collector in collectors:
        try:
            items = await collector.search(topic, limit)
            all_items.extend(items)
            results['sources'][name] = len(items)
        except Exception as e:
            results['sources'][name] = f"错误: {str(e)}"
        finally:
            await collector.close()

    # 分析
    synthesizer = ResearchSynthesizer()
    report = synthesizer.synthesize(all_items, topic)

    results['items'] = all_items
    results['report'] = {
        'topic': report.topic,
        'summary': report.summary,
        'key_findings': report.key_findings,
        'recommendations': report.recommendations,
        'created_at': report.created_at
    }

    return results


async def run_research_with_groq(topic: str, sources: list, limit: int):
    """执行研究 + Groq 深度分析"""
    # 先执行基础研究
    results = await run_research(topic, sources, limit)
    
    # 添加 Groq 深度分析
    if results['items']:
        groq_result = analyze_with_groq(results['items'], topic)
        results['groq_analysis'] = groq_result
    
    return results


@app.route('/')
def index():
    """首页"""
    return render_template('index.html')


@app.route('/research', methods=['POST'])
def research():
    """处理用户研究请求"""
    topic = request.form.get('topic', '').strip()
    sources = request.form.getlist('sources')
    
    if not topic:
        return render_template('index.html', error='请输入研究主题')
    
    if not sources:
        sources = ['hn', 'web']
    
    # 执行研究
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    results = loop.run_until_complete(run_research_with_groq(topic, sources, 15))
    loop.close()
    
    # 保存最新报告
    global latest_report
    latest_report = {
        'topic': topic,
        'report': results['report'],
        'groq_analysis': results.get('groq_analysis', {}),
        'items': results['items'],
        'created_at': datetime.now().isoformat()
    }
    
    return render_template('results.html',
                          topic=topic,
                          items=results['items'][:10],
                          report=results['report'],
                          groq_analysis=results.get('groq_analysis', {}))


@app.route('/api/research', methods=['POST'])
def api_research():
    """API: 执行研究"""
    data = request.get_json()
    topic = data.get('topic', '').strip()
    sources = data.get('sources', ['hn', 'web'])
    limit = int(data.get('limit', 10))

    if not topic:
        return jsonify({'error': '请输入研究主题'}), 400

    # 执行研究
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    results = loop.run_until_complete(run_research(topic, sources, limit))
    loop.close()

    # 保存历史
    history.append({
        'topic': topic,
        'time': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'count': len(results['items'])
    })

    return jsonify(results)


@app.route('/api/trigger', methods=['POST'])
def api_trigger():
    """API: 定时任务触发器（需Token验证）"""
    # Token 验证
    auth_header = request.headers.get('Authorization', '')
    token = auth_header.replace('Bearer ', '') if auth_header.startswith('Bearer ') else ''
    
    if token != API_TOKEN:
        return jsonify({'error': 'Unauthorized'}), 401
    
    # 执行研究（带 Groq 分析）
    topic = request.json.get('topic', 'AI前沿应用') if request.is_json else 'AI前沿应用'
    sources = ['hn', 'web']
    limit = 15
    
    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    results = loop.run_until_complete(run_research_with_groq(topic, sources, limit))
    loop.close()
    
    # 保存历史
    history.append({
        'topic': topic,
        'time': datetime.now().strftime('%Y-%m-%d %H:%M'),
        'count': len(results['items'])
    })
    
    # 保存最新报告（用于 Web UI）
    global latest_report
    latest_report = {
        'topic': topic,
        'report': results['report'],
        'groq_analysis': results.get('groq_analysis', {}),
        'items': results['items'],
        'created_at': datetime.now().isoformat()
    }
    
    return jsonify({
        'success': True,
        'report': results['report'],
        'groq_analysis': results.get('groq_analysis', {}),
        'items_count': len(results['items']),
        'triggered_at': datetime.now().isoformat()
    })


@app.route('/api/history')
def api_history():
    """API: 获取历史"""
    return jsonify(history[-20:])


@app.route('/api/debug')
def api_debug():
    """API: 调试环境变量"""
    deepseek_key = os.environ.get("DEEPSEEK_API_KEY", "")
    return jsonify({
        "deepseek_key_set": bool(deepseek_key),
        "deepseek_key_length": len(deepseek_key),
        "deepseek_key_preview": deepseek_key[:10] + "..." if deepseek_key else "EMPTY",
        "groq_key_set": bool(os.environ.get("GROQ_API_KEY")),
        "api_token_set": bool(API_TOKEN),
        "all_env_keys": list(os.environ.keys()),
        "relevant_keys": {k: v[:20] + "..." if len(v) > 20 else v for k, v in os.environ.items() if "DEEP" in k or "GROQ" in k or "API" in k or "KEY" in k}
    })


@app.route('/api/token')
def api_token():
    """API: 获取 Token（仅本地开发）"""
    if os.environ.get('RAILWAY_ENVIRONMENT'):
        return jsonify({'error': 'Not available in production'}), 403
    return jsonify({'token': API_TOKEN})


@app.route('/report')
def report_page():
    """Web UI: 显示完整报告"""
    if not latest_report['report']:
        return render_template('report.html', 
                              error="暂无报告，请先触发研究任务",
                              report=None,
                              groq_analysis=None)
    
    return render_template('report.html',
                          topic=latest_report['topic'],
                          report=latest_report['report'],
                          groq_analysis=latest_report['groq_analysis'],
                          created_at=latest_report['created_at'])


@app.route('/api/export/markdown')
def export_markdown():
    """API: 导出 Markdown 格式"""
    if not latest_report['report']:
        return jsonify({'error': 'No report available'}), 404
    
    # 构建 Markdown 内容
    md_content = f"""# AI Research Scout 报告

## 主题
{latest_report['topic']}

## 生成时间
{latest_report['created_at']}

## 摘要
{latest_report['report'].get('summary', 'N/A')}

## 关键发现
"""
    for finding in latest_report['report'].get('key_findings', []):
        md_content += f"- {finding}\n"
    
    md_content += "\n## 建议\n"
    for rec in latest_report['report'].get('recommendations', []):
        md_content += f"- {rec}\n"
    
    # 添加 LLM 分析
    if latest_report['groq_analysis'] and latest_report['groq_analysis'].get('success'):
        md_content += f"\n## LLM 深度分析\n\n{latest_report['groq_analysis'].get('analysis', '')}\n"
    
    # 添加项目列表
    md_content += "\n## 项目列表\n\n"
    for i, item in enumerate(latest_report['items'][:20], 1):
        md_content += f"{i}. **{item.get('title', 'N/A')}**\n"
        if item.get('url'):
            md_content += f"   - URL: {item.get('url')}\n"
    
    response = make_response(md_content)
    response.headers['Content-Type'] = 'text/markdown; charset=utf-8'
    response.headers['Content-Disposition'] = f'attachment; filename=report_{datetime.now().strftime("%Y%m%d")}.md'
    return response


@app.route('/api/export/txt')
def export_txt():
    """API: 导出 TXT 格式"""
    if not latest_report['report']:
        return jsonify({'error': 'No report available'}), 404
    
    txt_content = generate_report_text()
    
    response = make_response(txt_content)
    response.headers['Content-Type'] = 'text/plain; charset=utf-8'
    response.headers['Content-Disposition'] = f'attachment; filename=report_{datetime.now().strftime("%Y%m%d")}.txt'
    return response


@app.route('/api/export/word')
def export_word():
    """API: 导出 Word 格式"""
    if not latest_report['report']:
        return jsonify({'error': 'No report available'}), 404
    
    if not DOCX_AVAILABLE:
        return jsonify({'error': 'Word export not available. Install python-docx: pip install python-docx'}), 503
    
    doc = Document()
    
    # 标题
    title = doc.add_heading('AI Research Scout 报告', 0)
    
    # 主题
    doc.add_heading(f'主题: {latest_report["topic"]}', level=1)
    doc.add_paragraph(f'生成时间: {latest_report["created_at"]}')
    
    # 摘要
    doc.add_heading('摘要', level=2)
    doc.add_paragraph(latest_report['report'].get('summary', 'N/A'))
    
    # 关键发现
    doc.add_heading('关键发现', level=2)
    for finding in latest_report['report'].get('key_findings', []):
        doc.add_paragraph(finding, style='List Bullet')
    
    # 建议
    doc.add_heading('建议', level=2)
    for rec in latest_report['report'].get('recommendations', []):
        doc.add_paragraph(rec, style='List Bullet')
    
    # LLM 分析
    if latest_report['groq_analysis'] and latest_report['groq_analysis'].get('success'):
        doc.add_heading('LLM 深度分析', level=2)
        doc.add_paragraph(latest_report['groq_analysis'].get('analysis', ''))
    
    # 保存到内存
    doc_io = io.BytesIO()
    doc.save(doc_io)
    doc_io.seek(0)
    
    return send_file(
        doc_io,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=f'report_{datetime.now().strftime("%Y%m%d")}.docx'
    )


@app.route('/api/export/pdf')
def export_pdf():
    """API: 导出 PDF 格式"""
    if not latest_report['report']:
        return jsonify({'error': 'No report available'}), 404
    
    if not PDF_AVAILABLE:
        return jsonify({'error': 'PDF export not available. Install reportlab: pip install reportlab'}), 503
    
    pdf_io = io.BytesIO()
    doc = SimpleDocTemplate(pdf_io, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []
    
    # 标题
    story.append(Paragraph('AI Research Scout 报告', styles['Title']))
    story.append(Spacer(1, 12))
    
    # 主题
    story.append(Paragraph(f'主题: {latest_report["topic"]}', styles['Heading1']))
    story.append(Paragraph(f'生成时间: {latest_report["created_at"]}', styles['Normal']))
    story.append(Spacer(1, 12))
    
    # 摘要
    story.append(Paragraph('摘要', styles['Heading2']))
    story.append(Paragraph(latest_report['report'].get('summary', 'N/A'), styles['Normal']))
    story.append(Spacer(1, 12))
    
    # 关键发现
    story.append(Paragraph('关键发现', styles['Heading2']))
    for finding in latest_report['report'].get('key_findings', []):
        story.append(Paragraph(f'• {finding}', styles['Normal']))
    story.append(Spacer(1, 12))
    
    # 建议
    story.append(Paragraph('建议', styles['Heading2']))
    for rec in latest_report['report'].get('recommendations', []):
        story.append(Paragraph(f'• {rec}', styles['Normal']))
    
    doc.build(story)
    pdf_io.seek(0)
    
    return send_file(
        pdf_io,
        mimetype='application/pdf',
        as_attachment=True,
        download_name=f'report_{datetime.now().strftime("%Y%m%d")}.pdf'
    )


def generate_report_text():
    """生成报告文本"""
    txt_content = f"""AI Research Scout 报告
主题: {latest_report['topic']}
生成时间: {latest_report['created_at']}

摘要:
{latest_report['report'].get('summary', 'N/A')}

关键发现:
"""
    for finding in latest_report['report'].get('key_findings', []):
        txt_content += f"- {finding}\n"
    
    txt_content += "\n建议:\n"
    for rec in latest_report['report'].get('recommendations', []):
        txt_content += f"- {rec}\n"
    
    if latest_report['groq_analysis'] and latest_report['groq_analysis'].get('success'):
        txt_content += f"\n\nLLM 深度分析:\n{latest_report['groq_analysis'].get('analysis', '')}\n"
    
    return txt_content


if __name__ == '__main__':
    # 创建模板目录
    os.makedirs('templates', exist_ok=True)
    os.makedirs('static', exist_ok=True)
    
    # Railway 需要从环境变量获取端口
    port = int(os.environ.get('PORT', 5000))
    
    print("🚀 AI Research Scout Web UI")
    print(f"📍 Running on port: {port}")
    print(f"🔑 API Token: {API_TOKEN}")
    
    # Railway 需要 host='0.0.0.0'
    app.run(host='0.0.0.0', port=port, debug=False)